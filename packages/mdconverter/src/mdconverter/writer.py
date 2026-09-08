"""Long-form document stitching, outline chunking, and DOCX generation engine.

Provides deterministic tools for:
- Splitting long-form document outlines into manageable generation chunks.
- Stitching section responses into a unified, coherent Markdown document.
- Exporting long-form Markdown to DOCX with standard typography.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt


def chunk_outline_sections(outline_md: str) -> list[dict[str, Any]]:
    """Decompose an outline markdown into sequential chapter/section specs.

    Args:
        outline_md: Markdown text containing outline structure (#, ##, ###).

    Returns:
        List of section dictionaries containing section_id, level, title, and prompt context.
    """
    sections: list[dict[str, Any]] = []
    lines = outline_md.splitlines()

    current_section: dict[str, Any] | None = None
    section_index = 0

    for line in lines:
        stripped = line.strip()
        if not stripped:
            if current_section is not None and "description" in current_section:
                current_section["description"] += "\n"
            continue

        match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if match:
            if current_section is not None:
                current_section["description"] = current_section["description"].strip()
                sections.append(current_section)

            level = len(match.group(1))
            title = match.group(2).strip()
            section_index += 1
            current_section = {
                "section_id": f"sec_{section_index:02d}",
                "level": level,
                "title": title,
                "description": "",
            }
        else:
            if current_section is not None:
                current_section["description"] += stripped + "\n"

    if current_section is not None:
        current_section["description"] = current_section["description"].strip()
        sections.append(current_section)

    return sections


def stitch_markdown_sections(
    sections: list[dict[str, str]] | list[str],
    title: str | None = None,
) -> str:
    """Stitch multiple section contents into a single unified Markdown document.

    Args:
        sections: List of section strings or dictionaries with 'title' and 'content'.
        title: Optional document title.

    Returns:
        Consolidated markdown text string.
    """
    parts: list[str] = []
    if title:
        parts.append(f"# {title}\n")

    for sec in sections:
        if isinstance(sec, dict):
            sec_title = sec.get("title", "")
            content = sec.get("content", "").strip()
            level = sec.get("level", 2)
            header_prefix = "#" * int(level)
            if sec_title:
                parts.append(f"{header_prefix} {sec_title}\n\n{content}\n")
            else:
                parts.append(f"{content}\n")
        else:
            parts.append(f"{str(sec).strip()}\n")

    return "\n".join(parts)


def save_markdown_to_docx(text: str, filename: Path | str = "output.docx") -> Path:
    """Save markdown text directly to a DOCX document with Vietnamese font defaults.

    Args:
        text: Markdown content.
        filename: Target DOCX file path.

    Returns:
        Path to saved DOCX file.
    """
    out_path = Path(filename)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    for line in text.splitlines():
        trimmed = line.strip()
        if not trimmed:
            continue
        if trimmed.startswith("# "):
            doc.add_heading(trimmed[2:], level=1)
        elif trimmed.startswith("## "):
            doc.add_heading(trimmed[3:], level=2)
        elif trimmed.startswith("### "):
            doc.add_heading(trimmed[4:], level=3)
        elif trimmed.startswith("#### "):
            doc.add_heading(trimmed[5:], level=4)
        else:
            doc.add_paragraph(trimmed)

    doc.save(str(out_path))
    return out_path


__all__ = [
    "chunk_outline_sections",
    "stitch_markdown_sections",
    "save_markdown_to_docx",
]
