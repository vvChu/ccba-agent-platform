"""Tests for Phase 1 Deep Seams in mdconverter."""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from docx import Document

from mdconverter import (
    analyze_style_metrics,
    audit_microstructure,
    chunk_outline_sections,
    export_paper_to_docx,
    extract_docx_tables,
    format_all_qcvn_md_tables,
    format_qcvn_md_table,
    normalize_docx_markdown,
    parse_template_placeholders,
    redact_sensitive_info,
    save_markdown_to_docx,
    scaffold_manuscript,
    stitch_markdown_sections,
)

pytestmark = [pytest.mark.fast, pytest.mark.unit]


def test_audit_microstructure_cars_and_style():
    """Test academic microstructure audit for CARS moves and style warnings."""
    sample_paper = """---
title: "BIM Compliance Automated Checking"
---

# Introduction
The application of modern engineering standards has been widely studied in construction.
However, existing methods remain unclear regarding automated compliance verification.
In this paper, we propose an automated pipeline to bridge this critical gap.

# Materials and Methods
The proposed algorithm was implemented and evaluated using real-world IFC building models.
A total of 45 models were tested.

# Results
The empirical evaluation results demonstrate significant accuracy gains.

# Discussion
Our findings show high precision across complex test cases. Clearly, the framework is superior.
"""
    report = audit_microstructure(sample_paper)
    assert report.total_paragraphs >= 4
    assert "Introduction" in report.sections_found
    assert "Methods" in report.sections_found
    assert len(report.cars_moves_found) >= 2

    # Check for intensifier warning ("Clearly")
    findings = [f.message for f in report.findings if f.category == "Intensifier"]
    assert any("clearly" in m.lower() for m in findings)


def test_scaffold_manuscript():
    """Test generating academic manuscript skeleton."""
    with tempfile.TemporaryDirectory() as temp_dir:
        out_file = Path(temp_dir) / "paper_skeleton.md"
        scaffold_manuscript(out_file, title="Test Academic Paper")

        assert out_file.exists()
        content = out_file.read_text(encoding="utf-8")
        assert 'title: "Test Academic Paper"' in content
        assert "# Abstract" in content
        assert "# Introduction" in content
        assert "# Materials and Methods" in content
        assert "# References" in content


def test_export_paper_to_docx():
    """Test compiling academic markdown to standard DOCX."""
    with tempfile.TemporaryDirectory() as temp_dir:
        md_file = Path(temp_dir) / "source_paper.md"
        out_docx = Path(temp_dir) / "output_paper.docx"

        md_file.write_text(
            '---\ntitle: "Sample Title"\nauthors:\n  - name: "Author One"\n---\n\n# Introduction\nSample body text.',
            encoding="utf-8",
        )

        res = export_paper_to_docx(md_file, out_docx)
        assert res.exists()

        doc = Document(str(res))
        assert len(doc.paragraphs) >= 2
        assert any("Sample Title" in p.text for p in doc.paragraphs)
        assert any("Author One" in p.text for p in doc.paragraphs)


def test_table_extraction_and_qcvn_formatting():
    """Test extracting tables from DOCX and formatting 2D GFM tables."""
    with tempfile.TemporaryDirectory() as temp_dir:
        docx_path = Path(temp_dir) / "sample_tables.docx"
        doc = Document()
        doc.add_paragraph("Bảng 1 - Thông số kỹ thuật PCCC")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Tiêu chí"
        table.cell(0, 1).text = "Giá trị"
        table.cell(1, 0).text = "Bậc chịu lửa"
        table.cell(1, 1).text = "Bậc I"
        doc.save(str(docx_path))

        tables = extract_docx_tables(docx_path)
        assert len(tables) == 1
        t0 = tables[0]
        assert "bang_01" in t0.table_id or "bang_1" in t0.table_id
        assert t0.cols == 2
        assert "| Tiêu chí | Giá trị |" in t0.markdown_representation


def test_format_qcvn_md_table_raw():
    """Test normalizing raw pipe/tab text into clean GFM Markdown pipe table."""
    raw_table = "Cột A\tCột B\tCột C\nDòng 1\tDòng 2\tDòng 3"
    result = format_qcvn_md_table(raw_table)
    assert "| Cột A | Cột B | Cột C |" in result
    assert "| --- | --- | --- |" in result
    assert "| Dòng 1 | Dòng 2 | Dòng 3 |" in result


def test_format_all_qcvn_md_tables_in_file():
    """Test formatting broken table blocks across a markdown document."""
    with tempfile.TemporaryDirectory() as temp_dir:
        md_path = Path(temp_dir) / "qcvn_test.md"
        broken_content = """# Quy chuẩn kỹ thuật

### Bảng 2 - Bán kính phục vụ của trụ nước chữa cháy
Cột 1\tCột 2
Giá trị A\tGiá trị B
Giá trị C\tGiá trị D
CHÚ THÍCH: Trụ nước phải được bố trí dọc theo đường giao thông.

### Mục 2.2 Quy định chung
Nội dung điều khoản tiếp theo...
"""
        md_path.write_text(broken_content, encoding="utf-8")
        count = format_all_qcvn_md_tables(md_path)
        assert count == 1

        new_content = md_path.read_text(encoding="utf-8")
        assert "| Cột 1 | Cột 2 |" in new_content
        assert "_CHÚ THÍCH: Trụ nước phải được bố trí dọc theo đường giao thông._" in new_content


def test_normalize_docx_markdown():
    """Test mammoth converted markdown normalization."""
    raw_md = r"Điều 1\.1\-\(a\) __1.2 Tiêu chuẩn áp dụng__ \n\n### ### Bảng 1 - Phân loại"
    normalized = normalize_docx_markdown(raw_md)
    assert r"\." not in normalized
    assert r"\-" not in normalized
    assert r"\(" not in normalized
    assert "### 1.2 Tiêu chuẩn áp dụng" in normalized


def test_style_sanitization_and_metrics():
    """Test sensitive data redaction and style metric computation."""
    sensitive_text = "Liên hệ sk-abcdef1234567890abcdef1234567890 hoặc admin@example.com, ĐT: 0912345678, CCCD: 001090123456"
    redacted = redact_sensitive_info(sensitive_text)
    assert "sk-" not in redacted
    assert "[REDACTED_API_KEY]" in redacted
    assert "[REDACTED_EMAIL]" in redacted
    assert "[REDACTED_PHONE]" in redacted
    assert "[REDACTED_CCCD]" in redacted

    metrics = analyze_style_metrics(
        "Đây là một câu văn mẫu. Câu văn thứ hai dùng để kiểm tra độ dài trung bình."
    )
    assert metrics["total_sentences"] == 2
    assert metrics["total_words"] > 10
    assert metrics["avg_sentence_length"] > 0

    placeholders = parse_template_placeholders(
        "Kính gửi {{client_name}}, dự án {{project_name}} tại {{location}}."
    )
    assert set(placeholders) == {"client_name", "project_name", "location"}


def test_long_form_writer_tools():
    """Test outline chunking, section stitching, and docx saving."""
    outline = """# Title
## Section 1: Overview
Detail for overview.
## Section 2: Methodology
Detail for methods.
"""
    chunks = chunk_outline_sections(outline)
    assert len(chunks) == 3
    assert chunks[0]["title"] == "Title"
    assert chunks[1]["title"] == "Section 1: Overview"

    stitched = stitch_markdown_sections(
        [
            {"title": "Intro", "content": "Intro paragraph", "level": "2"},
            {"title": "Conclusion", "content": "Conclusion paragraph", "level": "2"},
        ],
        title="Document Master",
    )
    assert "# Document Master" in stitched
    assert "## Intro" in stitched
    assert "## Conclusion" in stitched

    with tempfile.TemporaryDirectory() as temp_dir:
        out_doc = Path(temp_dir) / "longform.docx"
        save_markdown_to_docx(stitched, out_doc)
        assert out_doc.exists()
        d = Document(str(out_doc))
        assert len(d.paragraphs) >= 3
