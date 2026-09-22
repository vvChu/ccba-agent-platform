"""Helper utilities to clean up text, process DOCX tables, and convert Markdown documents.

Includes LLM output parsing, OCR artifact removal, DOCX table reconstruction,
and Markdown-to-DOCX conversion.

Created by CCBA — Trung tâm Tư vấn và Ứng dụng BIM trong Xây dựng.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, TypeVar

from ccba_ai import parse_llm_json, strip_think_tags
from ccba_legal.normalizers import (
    NormalizerConfig,
    TextNormalizer,
    fix_common_ocr_typos,
    normalize_ocr_spacing,
)

try:
    import docx
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    DOCX_AVAILABLE = True
except ImportError:
    docx = None  # type: ignore[assignment]
    Document = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment,misc]
    Pt = None  # type: ignore[assignment,misc]
    DocxTable = None  # type: ignore[assignment,misc]
    DocxParagraph = None  # type: ignore[assignment,misc]
    DOCX_AVAILABLE = False

T = TypeVar("T")


class Cleaners:
    """Helper utilities to clean up text, extract JSON, and process DOCX documents."""

    @classmethod
    def strip_think_tags(cls, text: str) -> str:
        """Strip <think>...</think> tags and their contents from reasoning models."""
        return strip_think_tags(text)

    @classmethod
    def extract_json(
        cls,
        raw: str,
        schema: type[T] | None = None,
        strict: bool = False,
    ) -> Any:
        """Extract JSON dictionary or list from raw text containing Markdown fences."""
        return parse_llm_json(raw, schema=schema, strict=strict)

    @classmethod
    def remove_ocr_artifacts(cls, text: str) -> str:
        """Remove long uppercase lines commonly created by page headers/footers in OCR and fix broken spacing."""
        text = normalize_ocr_spacing(text)
        text = fix_common_ocr_typos(text)
        return re.sub(r"^[A-ZÀ-Ỹ][A-ZÀ-Ỹ\s_]{14,}\.?\s*$", "", text, flags=re.MULTILINE).strip()

    @classmethod
    def normalize_legal_text(
        cls,
        text: str,
        config: NormalizerConfig | None = None,
        strip_doc_id_prefix: str = "",
    ) -> str:
        """Apply full multi-pass legal OCR normalization chain."""
        normalizer = TextNormalizer(config=config)
        return normalizer.clean_chunk(text, strip_doc_id_prefix=strip_doc_id_prefix)

    @classmethod
    def convert_docx_table_to_markdown(cls, table: Any) -> str:
        """Convert a python-docx Table object to clean Markdown format."""
        if not hasattr(table, "rows"):
            return ""

        rows_data: list[list[str]] = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", "<br>") for cell in row.cells]
            rows_data.append(row_cells)

        if not rows_data or not rows_data[0]:
            return ""

        # Generate header
        header = "| " + " | ".join(rows_data[0]) + " |"
        separator = "| " + " | ".join(["---"] * len(rows_data[0])) + " |"

        body_rows = []
        for r in rows_data[1:]:
            # Handle cell row length mismatch
            if len(r) < len(rows_data[0]):
                r = r + [""] * (len(rows_data[0]) - len(r))
            elif len(r) > len(rows_data[0]):
                r = r[: len(rows_data[0])]
            body_rows.append("| " + " | ".join(r) + " |")

        return "\n".join([header, separator] + body_rows)

    @classmethod
    def extract_docx_with_tables(cls, docx_path: str | Path) -> str:
        """Extract clean text and reconstructed Markdown tables from a .docx file."""
        path = Path(docx_path)
        if not DOCX_AVAILABLE or not path.exists() or Document is None:
            return ""

        try:
            doc = Document(str(path))
            output_parts: list[str] = []

            for elem in doc.element.body:
                if elem.tag.endswith("p") and DocxParagraph is not None:
                    para = DocxParagraph(elem, doc)
                    txt = para.text.strip()
                    if txt:
                        output_parts.append(txt)
                elif elem.tag.endswith("table") and DocxTable is not None:
                    tbl = DocxTable(elem, doc)
                    md_tbl = cls.convert_docx_table_to_markdown(tbl)
                    if md_tbl:
                        output_parts.append("\n" + md_tbl + "\n")

            return "\n\n".join(output_parts)
        except Exception:
            return ""

    @classmethod
    def convert_markdown_to_docx(cls, md_path: str | Path, docx_path: str | Path) -> None:
        """Convert standard Markdown file to styled DOCX document (Times New Roman 12pt)."""
        if not DOCX_AVAILABLE or Document is None or Pt is None:
            raise RuntimeError(
                "python-docx is not installed or available in the current environment."
            )

        in_path = Path(md_path)
        out_path = Path(docx_path)

        if not in_path.exists():
            raise FileNotFoundError(f"Input file '{in_path}' does not exist.")

        document = Document()
        style = document.styles["Normal"]
        font = style.font
        font.name = "Times New Roman"
        font.size = Pt(12)

        lines = in_path.read_text(encoding="utf-8").splitlines()
        in_frontmatter = False

        for line in lines:
            line_str = line.strip()

            # Skip YAML Frontmatter
            if line_str == "---":
                in_frontmatter = not in_frontmatter
                continue
            if in_frontmatter or not line_str:
                continue

            # Handle Headers
            if line_str.startswith("# "):
                p = document.add_heading(line_str[2:], level=1)
                if WD_ALIGN_PARAGRAPH is not None:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if p.runs:
                    run = p.runs[0]
                    run.font.name = "Times New Roman"
                    run.font.bold = True
                    run.font.color.rgb = None
                continue

            if line_str.startswith("## "):
                p = document.add_heading(line_str[3:], level=2)
                if p.runs:
                    run = p.runs[0]
                    run.font.name = "Times New Roman"
                    run.font.bold = True
                    run.font.color.rgb = None
                continue

            if line_str.startswith("### "):
                p = document.add_heading(line_str[4:], level=3)
                if p.runs:
                    run = p.runs[0]
                    run.font.name = "Times New Roman"
                    run.font.bold = True
                    run.font.color.rgb = None
                continue

            # Centered bold titles
            if line_str.startswith("**") and line_str.endswith("**"):
                content = line_str[2:-2]
                if content.isupper() or "Độc lập - Tự do" in content:
                    p = document.add_paragraph()
                    if WD_ALIGN_PARAGRAPH is not None:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(content)
                    run.bold = True
                    continue

            # Lists
            if line_str.startswith("- ") or line_str.startswith("* "):
                document.add_paragraph(line_str[2:], style="List Bullet")
                continue

            if re.match(r"^\d+\.", line_str):
                content = re.sub(r"^\d+\.\s*", "", line_str)
                document.add_paragraph(content, style="List Number")
                continue

            # Normal Paragraph with inline bold
            p = document.add_paragraph()
            parts = re.split(r"(\*\*.*?\*\*)", line_str)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

        out_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(out_path))

    @classmethod
    def strip_administrative_noise(cls, text: str) -> str:
        """Strip administrative header/footer noise (Quoc hieu, Tieu ngu, Noi nhan, Signatures)."""
        # Cut off trailing administrative signature blocks
        noi_nhan_split = re.split(
            r"(?:\n\s*__\*?\s*Nơi nhận\s*:|\n\s*\*+Nơi nhận\s*:|\n\s*Nơi nhận\s*:|\n\s*__KT\.\s+BỘ\s+TRƯỞNG|\n\s*KT\.\s+BỘ\s+TRƯỞNG\s*\n|\n\s*__BỘ\s+TRƯỞNG__|\n\s*__THỨ\s+TRƯỞNG__|\n\s*__CHỦ\s+TỊCH\s+QUỐC\s+HỘI|\n\s*CHỦ\s+TỊCH\s+QUỐC\s+HỘI\s*\n|\n\s*__TM\.\s+QUỐC\s+HỘI|\n\s*__TM\.\s+CHÍNH\s+PHỦ|\n\s*__THỦ\s+TƯỚNG__|\n\s*\*+Luật\s+này\s+được\s+Quốc\s+hội|\n\s*Luật\s+này\s+được\s+Quốc\s+hội)",
            text,
            flags=re.IGNORECASE,
        )
        body = noi_nhan_split[0].strip()

        # Remove header noise
        body = re.sub(
            r"(?:^|\n)\s*(?:CỘNG\s+HÒA\s+XÃ\s+HỘI\s+CHỦ\s+NGHĨA\s+VIỆT\s+NAM|Độc\s+lập\s*-\s*Tự\s+do\s*-\s*Hạnh\s+phúc).*?(?=\n\n|#)",
            "",
            body,
            flags=re.DOTALL | re.IGNORECASE,
        )
        return body.strip()

    @classmethod
    def strip_web_artifacts(cls, text: str) -> str:
        """Strip web scraping HTML tags and tracking scripts."""
        text = re.sub(r"<script.*?>.*?</script>", "", text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r"<form.*?>.*?</form>", "", text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r"<iframe.*?>.*?</iframe>", "", text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r"<input.*?>", "", text, flags=re.IGNORECASE)
        text = re.sub(r'class="NoiDungChiase"', "", text, flags=re.IGNORECASE)
        text = re.sub(r'onclick=".*?"', "", text, flags=re.IGNORECASE)
        return text


# Procedural convenience aliases
convert_docx_table_to_markdown = Cleaners.convert_docx_table_to_markdown
extract_docx_with_tables = Cleaners.extract_docx_with_tables
convert_markdown_to_docx = Cleaners.convert_markdown_to_docx
strip_think_tags_clean = Cleaners.strip_think_tags
extract_json_clean = Cleaners.extract_json
remove_ocr_artifacts = Cleaners.remove_ocr_artifacts
normalize_legal_text = Cleaners.normalize_legal_text
strip_administrative_noise = Cleaners.strip_administrative_noise
strip_web_artifacts = Cleaners.strip_web_artifacts
