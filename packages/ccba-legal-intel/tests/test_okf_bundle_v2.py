"""Unit tests for OKF Bundle v2.0 specification (ADR 0038) in ccba-legal-intel."""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

import pytest
import yaml

from ccba_legal.packager import OKFBundlePackager


def test_okf_v2_metadata_yaml_structure() -> None:
    """Test metadata.yaml creation and fields compliance with ADR 0038."""
    with tempfile.TemporaryDirectory() as tmpdir:
        root_dir = Path(tmpdir)
        packager = OKFBundlePackager(root_dir)

        metadata = {
            "title": "Nghị định quy định chi tiết một số điều của Luật Xây dựng",
            "type": "vbpl",
            "document_number": "217/2026/NĐ-CP",
            "issued_by": "Chính phủ",
            "issued_date": "2026-06-19",
            "effective_date": "2026-07-01",
            "source_url": "https://thuvienphapluat.vn/van-ban/217-2026-ND-CP",
            "sha256": "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
        }

        bundle_dir = packager.package_bundle_v2(
            doc_id="nghi_dinh_217_2026_nd_cp",
            content="# Nghị định 217\n\nNội dung văn bản...",
            metadata=metadata,
        )

        meta_file = bundle_dir / "metadata.yaml"
        assert meta_file.exists()

        meta = yaml.safe_load(meta_file.read_text(encoding="utf-8"))
        assert meta["doc_id"] == "nghi_dinh_217_2026_nd_cp"
        assert meta["doc_number"] == "217/2026/NĐ-CP"
        assert meta["title"] == "Nghị định quy định chi tiết một số điều của Luật Xây dựng"
        assert meta["type"] == "vbpl"
        assert meta["issuer"] == "Chính phủ"
        assert meta["issued_date"] == "2026-06-19"
        assert meta["effective_date"] == "2026-07-01"
        assert meta["status"] == "effective"
        assert meta["sha256"] == "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855"
        assert "timestamp" in meta


def test_okf_v2_clauses_json_hierarchy_and_lines() -> None:
    """Test AST clauses.json generation with flat index, hierarchy, and line coordinates."""
    sample_text = """# Nghị định 217/2026/NĐ-CP

## Chương I: Quy định chung
### Mục 1: Phạm vi và đối tượng
### Điều 1. Phạm vi điều chỉnh
1. Nghị định này quy định chi tiết thi hành Luật Xây dựng.
a) Quản lý dự án đầu tư xây dựng;
b) Thẩm định thiết kế và dự toán.
2. Áp dụng đối với cơ quan, tổ chức, cá nhân liên quan.

### Điều 2. Đối tượng áp dụng
1. Cơ quan quản lý nhà nước về xây dựng.
"""
    with tempfile.TemporaryDirectory() as tmpdir:
        root_dir = Path(tmpdir)
        packager = OKFBundlePackager(root_dir)

        clauses = packager.generate_clauses_json(sample_text, root_dir)
        assert len(clauses) == 9

        clauses_file = root_dir / "clauses.json"
        assert clauses_file.exists()

        loaded_clauses = json.loads(clauses_file.read_text(encoding="utf-8"))
        assert len(loaded_clauses) == 9

        # 1. Chapter node
        c1 = loaded_clauses[0]
        assert c1["clause_id"] == "chuong-i"
        assert c1["node_type"] == "chapter"
        assert c1["parent_id"] is None
        assert c1["line_start"] == 3

        # 2. Section node
        s1 = loaded_clauses[1]
        assert s1["clause_id"] == "muc-1"
        assert s1["node_type"] == "section"
        assert s1["parent_id"] == "chuong-i"
        assert s1["line_start"] == 4

        # 3. Article 1 node
        a1 = loaded_clauses[2]
        assert a1["clause_id"] == "dieu-1"
        assert a1["node_type"] == "article"
        assert a1["parent_id"] == "muc-1"
        assert a1["line_start"] == 5

        # 4. Clause 1 node
        k1 = loaded_clauses[3]
        assert k1["clause_id"] == "dieu-1-khoan-1"
        assert k1["node_type"] == "clause"
        assert k1["parent_id"] == "dieu-1"
        assert k1["line_start"] == 6

        # 5. Point a node
        pa = loaded_clauses[4]
        assert pa["clause_id"] == "dieu-1-khoan-1-diem-a"
        assert pa["node_type"] == "point"
        assert pa["parent_id"] == "dieu-1-khoan-1"
        assert pa["line_start"] == 7

        # 6. Point b node
        pb = loaded_clauses[5]
        assert pb["clause_id"] == "dieu-1-khoan-1-diem-b"
        assert pb["node_type"] == "point"
        assert pb["parent_id"] == "dieu-1-khoan-1"
        assert pb["line_start"] == 8

        # 7. Clause 2 node
        k2 = loaded_clauses[6]
        assert k2["clause_id"] == "dieu-1-khoan-2"
        assert k2["node_type"] == "clause"
        assert k2["parent_id"] == "dieu-1"
        assert k2["line_start"] == 9

        # 8. Article 2 node
        a2 = loaded_clauses[7]
        assert a2["clause_id"] == "dieu-2"
        assert a2["node_type"] == "article"
        assert a2["parent_id"] == "muc-1"
        assert a2["line_start"] == 11

        # 9. Clause 1 of Article 2
        a2_k1 = loaded_clauses[8]
        assert a2_k1["clause_id"] == "dieu-2-khoan-1"
        assert a2_k1["node_type"] == "clause"
        assert a2_k1["parent_id"] == "dieu-2"
        assert a2_k1["line_start"] == 12


def test_okf_v2_qa_benchmark_validation() -> None:
    """Test 5-field schema validation for qa_benchmark.json."""
    with tempfile.TemporaryDirectory() as tmpdir:
        root_dir = Path(tmpdir)
        packager = OKFBundlePackager(root_dir)

        valid_qa = [
            {
                "question": "Cơ quan nào có thẩm quyền thẩm định Báo cáo nghiên cứu khả thi dự án nhóm A?",
                "answer": "Cơ quan chuyên môn về xây dựng thuộc Bộ quản lý công trình xây dựng chuyên ngành.",
                "anchor": "dieu-14-khoan-2",
                "citation": "Khoản 2 Điều 14 Nghị định 217/2026/NĐ-CP",
                "ground_truth_context": "2. Đối với dự án nhóm A do Thủ tướng Chính phủ quyết định chủ trương đầu tư...",
            }
        ]

        out_path = packager.write_qa_benchmark(valid_qa, root_dir)
        assert out_path.exists()

        data = json.loads(out_path.read_text(encoding="utf-8"))
        assert len(data) == 1
        assert data[0]["citation"] == "Khoản 2 Điều 14 Nghị định 217/2026/NĐ-CP"

        # Test invalid QA missing ground_truth_context
        invalid_qa = [
            {
                "question": "Câu hỏi mẫu?",
                "answer": "Câu trả lời mẫu",
                "anchor": "dieu-1",
                "citation": "Điều 1",
                # missing ground_truth_context
            }
        ]

        with pytest.raises(ValueError, match="missing or empty required fields"):
            packager.write_qa_benchmark(invalid_qa, root_dir)


def test_okf_v2_table_reconstructor_integration() -> None:
    """Test integrating TableReconstructor to extract tables and save to tables/json and tables/csv."""
    with tempfile.TemporaryDirectory() as tmpdir:
        root_dir = Path(tmpdir)
        packager = OKFBundlePackager(root_dir)

        # Create a mock docx if docx is installed
        try:
            from docx import Document

            docx_file = root_dir / "sample.docx"
            doc = Document()
            doc.add_paragraph("Bảng 1 - Phân loại cấp công trình")
            table = doc.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "Cấp công trình"
            table.rows[0].cells[1].text = "Quy mô"
            table.rows[1].cells[0].text = "Cấp I"
            table.rows[1].cells[1].text = "> 100m"
            doc.save(str(docx_file))

            tables = packager.integrate_tables(docx_file, root_dir)
            assert len(tables) == 1

            slug = tables[0].table_id
            assert slug == "bang_01_phan_loai_cap_cong_trinh"
            tables_json = root_dir / "tables" / "json" / f"{slug}.json"
            tables_csv = root_dir / "tables" / "csv" / f"{slug}.csv"

            assert tables_json.exists()
            assert tables_csv.exists()
        except ImportError:
            pytest.skip("python-docx not installed")


def test_okf_v2_package_bundle_full_pipeline() -> None:
    """Test full package_bundle_v2 pipeline producing all ADR 0038 artifacts."""
    with tempfile.TemporaryDirectory() as tmpdir:
        root_dir = Path(tmpdir)
        packager = OKFBundlePackager(root_dir)

        content = """# Nghị định 217/2026/NĐ-CP

## Chương I: Quy định chung
### Điều 1. Phạm vi điều chỉnh
1. Quy định về quản lý hoạt động xây dựng.
"""
        metadata = {
            "title": "Nghị định 217/2026/NĐ-CP",
            "type": "vbpl",
            "document_number": "217/2026/NĐ-CP",
            "issued_by": "Chính phủ",
            "issued_date": "2026-06-19",
            "effective_date": "2026-07-01",
        }
        qa_items = [
            {
                "question": "Phạm vi điều chỉnh của Nghị định 217 là gì?",
                "answer": "Quy định về quản lý hoạt động xây dựng.",
                "anchor": "dieu-1-khoan-1",
                "citation": "Khoản 1 Điều 1 Nghị định 217/2026/NĐ-CP",
                "ground_truth_context": "1. Quy định về quản lý hoạt động xây dựng.",
            }
        ]

        bundle_dir = packager.package_bundle_v2(
            doc_id="nghi_dinh_217_2026_nd_cp",
            content=content,
            metadata=metadata,
            qa_items=qa_items,
        )

        assert bundle_dir.exists()
        assert (bundle_dir / "metadata.yaml").exists()
        assert (bundle_dir / "nghi_dinh_217_2026_nd_cp.md").exists()
        assert (bundle_dir / "clauses.json").exists()
        assert (bundle_dir / "qa_benchmark.json").exists()
        assert (bundle_dir / "index.md").exists()
        assert (bundle_dir / "log.md").exists()

        # Check pure markdown content (zero frontmatter)
        md_text = (bundle_dir / "nghi_dinh_217_2026_nd_cp.md").read_text(encoding="utf-8")
        assert not md_text.startswith("---")
        assert "# Nghị định 217/2026/NĐ-CP" in md_text
