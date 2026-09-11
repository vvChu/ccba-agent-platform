# Copyright (c) 2026 CCBA. All rights reserved.
"""Unit tests for Virtual 2D Grid Engine, ADR 0041 Table Knowledge Extraction, and Pipe Escaping."""

from ccba_legal.converters.standard.handlers.table_handler import (
    build_composite_headers,
    detect_table_archetype,
    escape_table_pipes,
    resolve_hierarchical_headers,
)


def test_escape_table_pipes_in_prose_and_math():
    """Verify that unescaped pipe symbols in prose and math expressions are safely escaped."""
    # Prose pipe
    prose = "Cột A | Cột B"
    assert escape_table_pipes(prose) == r"Cột A \| Cột B"

    # Math absolute value
    math_abs = r"Điều kiện $|x| \le 1$ và $y > 0$"
    escaped_math = escape_table_pipes(math_abs)
    assert r"\vert " in escaped_math
    assert r"|" not in escaped_math.replace(r"\|", "").replace(r"\vert", "")

    # Conditional probability or set notation
    math_set = r"Tập hợp $\{x | x \ge 0\}$"
    escaped_set = escape_table_pipes(math_set)
    assert r"\vert " in escaped_set


def test_build_composite_headers():
    """Verify combining multi-row table headers into single composite headers."""
    headers = [
        ["Thông số", "Giới hạn cho phép", "Giới hạn cho phép"],
        ["—", "Tối thiểu", "Tối đa"],
    ]
    composite = build_composite_headers(headers)
    assert composite == ["Thông số", "Giới hạn cho phép — Tối thiểu", "Giới hạn cho phép — Tối đa"]


def test_resolve_hierarchical_headers():
    """Verify resolving hierarchical multi-tier headers into composite single row."""
    grid = [
        ["Chỉ tiêu", "Tầng 1", "Tầng 1"],
        ["—", "Mức A", "Mức B"],
        ["Độ bền", "100", "200"],
    ]
    resolved = resolve_hierarchical_headers(grid)
    assert resolved[0] == ["Chỉ tiêu", "Tầng 1 — Mức A", "Tầng 1 — Mức B"]
    assert resolved[1] == ["Độ bền", "100", "200"]


def test_detect_table_archetype():
    """Verify table archetype classification according to ADR 0041."""
    # Administrative form
    assert (
        detect_table_archetype(
            rows_count=10,
            cols_count=4,
            text="Biên bản kiểm tra Mẫu số 01 chức vụ của người ký",
        )
        == "ADMIN_FORM"
    )

    # Footnote rich table
    assert (
        detect_table_archetype(
            rows_count=5,
            cols_count=3,
            text="Bảng tra cứu kỹ thuật",
            has_footnotes=True,
        )
        == "FOOTNOTE_RICH"
    )

    # Borderless layout (signature or national emblem)
    assert (
        detect_table_archetype(
            rows_count=2,
            cols_count=2,
            text="Cộng hòa xã hội chủ nghĩa Việt Nam Nơi nhận:",
            is_captioned=False,
        )
        == "BORDERLESS_LAYOUT"
    )


def test_classify_and_extract_tables_with_oxml_table(tmp_path):
    """Verify that classify_and_extract_tables extracts tables from docx without xpath TypeError."""
    import docx

    from ccba_legal.converters.table_extractor import classify_and_extract_tables

    doc = docx.Document()
    doc.add_paragraph("Bảng 1 — Thông số tính toán thiết kế")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Chỉ tiêu"
    table.cell(0, 1).text = "Giá trị"
    table.cell(1, 0).text = "Độ bền kéo"
    table.cell(1, 1).text = "400 MPa"

    docx_path = tmp_path / "test.docx"
    doc.save(docx_path)
    bundle_dir = tmp_path / "bundle"
    catalog = classify_and_extract_tables(docx_path, bundle_dir)

    assert len(catalog) == 1
    assert catalog[0]["table_id"] == "bang_01"
    tables_dir = bundle_dir / "tables"
    assert (tables_dir / "csv" / f"{catalog[0]['table_id']}.csv").exists()
    assert (tables_dir / "json" / f"{catalog[0]['table_id']}.json").exists()
