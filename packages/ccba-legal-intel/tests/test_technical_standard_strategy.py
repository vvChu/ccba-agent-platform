# Copyright (c) 2026 CCBA. All rights reserved.
"""Unit tests for Modular Technical Standard Strategy Converter (OKF v2.4)."""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document

from ccba_legal.converters.standard.strategy import (
    StandardConversionContext,
    process_technical_standard_strategy,
)


def test_process_technical_standard_strategy_dummy_docx() -> None:
    """Test process_technical_standard_strategy with a minimal dummy technical standard DOCX."""
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        docx_file = tmp_path / "tcvn_9999_2026.docx"
        bundle_dir = tmp_path / "tcvn_9999_2026"

        # Create dummy DOCX
        doc = Document()
        doc.add_paragraph("TIÊU CHUẨN QUỐC GIA")
        doc.add_paragraph("TCVN 9999:2026")
        doc.add_paragraph("Lời nói đầu")
        doc.add_paragraph("TCVN 9999:2026 do Viện KHCN biên soạn.")
        doc.add_paragraph("1. QUY ĐỊNH CHUNG")
        doc.add_paragraph("1.1 Phạm vi áp dụng")
        doc.add_paragraph("Tiêu chuẩn này quy định các yêu cầu kỹ thuật.")

        # Formula paragraph
        doc.add_paragraph("RA x Ia <= 50 \t(1)")

        # Table caption & table
        doc.add_paragraph("Bảng 1 - Thông số kỹ thuật thiết kế")
        tbl = doc.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].text = "Thông số"
        tbl.rows[0].cells[1].text = "Giá trị"
        tbl.rows[1].cells[0].text = "Cường độ chịu kéo"
        tbl.rows[1].cells[1].text = "250 MPa"

        # Annex heading and content
        doc.add_paragraph("PHỤ LỤC A (Quy định) Phương pháp thử tải trọng")
        doc.add_paragraph("Thực hiện thử nghiệm theo quy trình chuẩn.")

        doc.save(str(docx_file))

        res = process_technical_standard_strategy(
            docx_path=docx_file,
            bundle_dir=bundle_dir,
            output_filename="tcvn_9999_2026.md",
            doc_meta={"id": "tcvn_9999_2026", "document_number": "TCVN 9999:2026"},
        )

        assert res["status"] == "success"
        assert res["archetype"] == "TECHNICAL_TCVN"
        assert res["tables_count"] == 1

        # Check bundle files
        md_file = bundle_dir / "tcvn_9999_2026.md"
        assert md_file.exists()
        md_text = md_file.read_text(encoding="utf-8")
        assert "okf_version: '2.4'" in md_text or 'okf_version: "2.4"' in md_text
        assert "## 1  QUY ĐỊNH CHUNG" in md_text
        assert "### 1.1  Phạm vi áp dụng" in md_text
        assert "formula-1" in md_text
        assert "DANH MỤC PHỤ LỤC KỸ THUẬT CHUYÊN ĐỀ" in md_text

        # Check tables catalog and files
        tables_cat_file = bundle_dir / "tables" / "tables_catalog.json"
        assert tables_cat_file.exists()
        csv_file = bundle_dir / "tables" / "csv" / "bang_01.csv"
        assert csv_file.exists()
        json_file = bundle_dir / "tables" / "json" / "bang_01.json"
        assert json_file.exists()

        # Check modular annex
        annexes = list((bundle_dir / "annexes").glob("*.md"))
        assert len(annexes) == 1
        assert "phu_luc_a" in annexes[0].name

        # Check metadata.yaml
        meta_file = bundle_dir / "metadata.yaml"
        assert meta_file.exists()


def test_standard_conversion_context_initialization() -> None:
    """Test StandardConversionContext properties and buffer management."""
    with tempfile.TemporaryDirectory() as tmpdir:
        bundle_dir = Path(tmpdir) / "test_bundle"
        ctx = StandardConversionContext(bundle_dir=bundle_dir, output_filename="test.md")

        assert ctx.current_target == "main"
        assert ctx.active_parts is ctx.body_md_parts

        ctx.emit("Hello World\n")
        assert ctx.body_md_parts == ["Hello World\n"]

        # Switch to annex target
        ctx.annex_buffers["A"] = {"parts": [], "title": "Phụ lục A", "slug": "phu_luc_a"}
        ctx.current_target = "A"
        assert ctx.active_parts is ctx.annex_buffers["A"]["parts"]
        ctx.emit("Annex content\n")
        assert ctx.annex_buffers["A"]["parts"] == ["Annex content\n"]


def test_process_technical_standard_strategy_multipart_disambiguation() -> None:
    """Test multi-part standard table disambiguation (ADR 0044) and footnote decoupling (ADR 0041)."""
    import json

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        docx_file = tmp_path / "qcvn_99_2026.docx"
        bundle_dir = tmp_path / "qcvn_99_2026"

        doc = Document()
        doc.add_paragraph("QUY CHUẨN KỸ THUẬT QUỐC GIA")
        doc.add_paragraph("QCVN 99:2026/BXD")
        doc.add_paragraph("Lời nói đầu")
        doc.add_paragraph("Quy chuẩn này gồm các phần độc lập.")

        # Part 1
        doc.add_paragraph("QCVN 99-1:2026/BXD")
        doc.add_paragraph("QUY CHUẨN KỸ THUẬT QUỐC GIA - PHẦN 1")
        doc.add_paragraph("1. QUY ĐỊNH CHUNG")
        doc.add_paragraph("1.1 Phạm vi áp dụng")
        doc.add_paragraph("Bảng 1 - Thông số Phần 1")
        tbl1 = doc.add_table(rows=2, cols=2)
        tbl1.rows[0].cells[0].text = "Chỉ tiêu"
        tbl1.rows[0].cells[1].text = "Trị số P1"
        tbl1.rows[1].cells[0].text = "Chỉ tiêu A"
        tbl1.rows[1].cells[1].text = "10"

        # Part 2
        doc.add_paragraph("QCVN 99-2:2026/BXD")
        doc.add_paragraph("QUY CHUẨN KỸ THUẬT QUỐC GIA - PHẦN 2")
        doc.add_paragraph("1. QUY ĐỊNH CHUNG")
        doc.add_paragraph("1.1 Phạm vi áp dụng")
        doc.add_paragraph("Bảng 1 - Thông số Phần 2")
        tbl2 = doc.add_table(rows=3, cols=2)
        tbl2.rows[0].cells[0].text = "Chỉ tiêu"
        tbl2.rows[0].cells[1].text = "Trị số P2"
        tbl2.rows[1].cells[0].text = "Chỉ tiêu B"
        tbl2.rows[1].cells[1].text = "20"
        # Footnote row in tbl2
        tbl2.rows[2].cells[0].text = "CHÚ THÍCH: Ghi chú chân bảng phần 2"
        tbl2.rows[2].cells[1].text = "CHÚ THÍCH: Ghi chú chân bảng phần 2"

        doc.save(str(docx_file))

        res = process_technical_standard_strategy(
            docx_path=docx_file,
            bundle_dir=bundle_dir,
            output_filename="qcvn_99_2026.md",
            doc_meta={"id": "qcvn_99_2026", "document_number": "QCVN 99:2026/BXD"},
        )

        assert res["status"] == "success"
        assert res["tables_count"] == 2

        # Check catalog
        cat_file = bundle_dir / "tables" / "tables_catalog.json"
        assert cat_file.exists()
        cat_data = json.loads(cat_file.read_text(encoding="utf-8"))
        tables = cat_data["tables"]
        assert len(tables) == 2

        # Disambiguated table IDs
        t_ids = [t["table_id"] for t in tables]
        assert t_ids == ["bang_p01_01", "bang_p02_01"]
        assert tables[0]["part_id"] == "p01"
        assert tables[1]["part_id"] == "p02"

        # Files exist and are distinct
        csv1 = bundle_dir / "tables" / "csv" / "bang_p01_01.csv"
        csv2 = bundle_dir / "tables" / "csv" / "bang_p02_01.csv"
        assert csv1.exists()
        assert csv2.exists()

        # Footnote decoupling: tbl2 footnote is NOT in CSV rows
        csv2_content = csv2.read_text(encoding="utf-8")
        assert "CHÚ THÍCH:" not in csv2_content
        assert "Chỉ tiêu B" in csv2_content

        # Footnote is in json metadata
        json2 = json.loads(
            (bundle_dir / "tables" / "json" / "bang_p02_01.json").read_text(encoding="utf-8")
        )
        assert json2["part_id"] == "p02"
        assert len(json2["footnotes"]) > 0


def test_process_technical_standard_strategy_roman_numerals_and_superscript_footnotes() -> None:
    """Test multi-part standard with Roman numerals (PHẦN I, PHẦN II) and superscript footnotes."""
    import json

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        docx_file = tmp_path / "tcvn_roman_2026.docx"
        bundle_dir = tmp_path / "tcvn_roman_2026"

        doc = Document()
        doc.add_paragraph("TIÊU CHUẨN QUỐC GIA")
        doc.add_paragraph("TCVN 8888:2026")
        doc.add_paragraph("Lời nói đầu")
        doc.add_paragraph("Tiêu chuẩn gồm các phần theo chữ số La Mã.")

        # Part I
        doc.add_paragraph("PHẦN I: YÊU CẦU CHUNG")
        doc.add_paragraph("1. QUY ĐỊNH CHUNG")
        doc.add_paragraph("1.1 Phạm vi")
        doc.add_paragraph("Bảng 1 - Chỉ tiêu phần I")
        tbl1 = doc.add_table(rows=3, cols=2)
        tbl1.rows[0].cells[0].text = "Thông số"
        tbl1.rows[0].cells[1].text = "Giá trị"
        tbl1.rows[1].cells[0].text = "Tham số 1"
        tbl1.rows[1].cells[1].text = "100"
        # Superscript footnote row in tbl1
        p_fn = tbl1.rows[2].cells[0].paragraphs[0]
        r1 = p_fn.add_run("1) ")
        r1.font.superscript = True
        p_fn.add_run("Ghi chú chân bảng số một")
        tbl1.rows[2].cells[1].text = ""

        # Part II
        doc.add_paragraph("PHẦN II: YÊU CẦU ĐẶC BIỆT")
        doc.add_paragraph("1. QUY ĐỊNH CHUNG")
        doc.add_paragraph("1.1 Phạm vi phần II")
        doc.add_paragraph("Bảng 1 - Chỉ tiêu phần II")
        tbl2 = doc.add_table(rows=2, cols=2)
        tbl2.rows[0].cells[0].text = "Thông số"
        tbl2.rows[0].cells[1].text = "Giá trị"
        tbl2.rows[1].cells[0].text = "Tham số 2"
        tbl2.rows[1].cells[1].text = "200"

        doc.save(str(docx_file))

        res = process_technical_standard_strategy(
            docx_path=docx_file,
            bundle_dir=bundle_dir,
            output_filename="tcvn_8888_2026.md",
            doc_meta={"id": "tcvn_8888_2026", "document_number": "TCVN 8888:2026"},
        )

        assert res["status"] == "success"
        assert res["tables_count"] == 2

        cat_file = bundle_dir / "tables" / "tables_catalog.json"
        cat_data = json.loads(cat_file.read_text(encoding="utf-8"))
        tables = cat_data["tables"]
        assert len(tables) == 2

        # Check Roman numeral disambiguation (p01, p02)
        assert tables[0]["table_id"] == "bang_p01_01"
        assert tables[0]["part_id"] == "p01"
        assert tables[1]["table_id"] == "bang_p02_01"
        assert tables[1]["part_id"] == "p02"

        # Check footnote decoupling: tbl1 footnote row NOT in CSV
        csv1 = (bundle_dir / "tables" / "csv" / "bang_p01_01.csv").read_text(encoding="utf-8")
        assert "Ghi chú chân bảng số một" not in csv1
        assert "Tham số 1" in csv1

        # Check footnote is in JSON metadata
        json1 = json.loads(
            (bundle_dir / "tables" / "json" / "bang_p01_01.json").read_text(encoding="utf-8")
        )
        assert len(json1["footnotes"]) > 0
