"""Tests for Phase 1 Deep Seams in ccba-ooxml."""

from __future__ import annotations

import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from pptx import Presentation

from ccba_ooxml import (
    FormattingProfile,
    convert_md_to_docx,
    find_soffice_bin,
    format_docx,
    generate_thumbnails,
    get_soffice_env,
    pptx_inventory,
    pptx_replace_text,
    rearrange_slides,
    run_soffice,
)

pytestmark = [pytest.mark.fast, pytest.mark.unit]


def test_formatting_profile_defaults():
    """Verify default formatting profile values matching administrative standards."""
    profile = FormattingProfile()
    assert profile.font_name == "Times New Roman"
    assert profile.font_size_pt == 13.0
    assert profile.line_spacing == 1.15
    assert profile.margin_top_cm == 2.5


def test_format_docx_existing_document():
    """Test formatting an existing document with custom profile."""
    with tempfile.TemporaryDirectory() as temp_dir:
        src_docx = Path(temp_dir) / "source.docx"
        out_docx = Path(temp_dir) / "reformatted.docx"

        doc = Document()
        doc.add_paragraph("Original paragraph text")
        doc.save(str(src_docx))

        custom_profile = FormattingProfile(font_name="Arial", font_size_pt=11.0)
        res = format_docx(input_path=src_docx, output_path=out_docx, profile=custom_profile)

        assert res.exists()
        reformatted_doc = Document(str(res))
        assert len(reformatted_doc.paragraphs) >= 1
        assert "Original paragraph text" in reformatted_doc.paragraphs[0].text


def test_convert_md_to_docx():
    """Test markdown to DOCX compilation seam."""
    with tempfile.TemporaryDirectory() as temp_dir:
        md_file = Path(temp_dir) / "input.md"
        md_file.write_text(
            "# Chapter 1: Introduction\n\nContent for chapter 1.\n\n## Section 1.1\nSubsection details.",
            encoding="utf-8",
        )
        out_docx = Path(temp_dir) / "compiled.docx"

        result = convert_md_to_docx(md_file, out_docx)
        assert result.exists()

        doc = Document(str(result))
        headings = [p.text for p in doc.paragraphs if p.text.startswith("Chapter")]
        assert len(headings) == 1


def test_soffice_helpers():
    """Test soffice environment helper functions."""
    env = get_soffice_env()
    assert isinstance(env, dict)
    assert "SAL_USE_VCLPLUGIN" in env
    assert env["SAL_USE_VCLPLUGIN"] == "svp"

    with patch("shutil.which", return_value="/usr/bin/soffice"):
        bin_path = find_soffice_bin()
        assert bin_path == "/usr/bin/soffice"

    with patch("subprocess.run") as mock_run:
        mock_run.return_value = MagicMock(returncode=0, stdout="converted", stderr="")
        with patch("ccba_ooxml.soffice.find_soffice_bin", return_value="soffice"):
            res = run_soffice(["--headless", "--convert-to", "pdf", "input.docx"])
            assert res.returncode == 0


def test_pptx_inventory():
    """Test PPTX inventory extraction seam."""
    with tempfile.TemporaryDirectory() as temp_dir:
        pptx_path = Path(temp_dir) / "test_presentation.pptx"

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        title.text = "Hello Phase 1 PPTX Seams"
        subtitle = slide.placeholders[1]
        subtitle.text = "Subtitle test content"
        prs.save(str(pptx_path))

        inventory = pptx_inventory(pptx_path)
        assert inventory.total_slides == 1
        assert len(inventory.slides) == 1
        slide_info = inventory.slides[0]
        assert slide_info["slide_id"] == "slide-0"
        assert len(slide_info["shapes"]) >= 1


def test_pptx_replace_text():
    """Test PPTX text replacement seam."""
    with tempfile.TemporaryDirectory() as temp_dir:
        src_pptx = Path(temp_dir) / "source.pptx"
        dst_pptx = Path(temp_dir) / "replaced.pptx"

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        title.text = "Template Project: {{project_name}}"
        prs.save(str(src_pptx))

        replacements = {"{{project_name}}": "National High-Tech Hospital"}
        res_path = pptx_replace_text(src_pptx, replacements, output_path=dst_pptx)

        assert res_path.exists()
        res_prs = Presentation(str(res_path))
        slide_0 = res_prs.slides[0]
        assert "National High-Tech Hospital" in slide_0.shapes.title.text


def test_pptx_rearrange_slides():
    """Test PPTX slide rearrangement seam."""
    with tempfile.TemporaryDirectory() as temp_dir:
        src_pptx = Path(temp_dir) / "multi_slides.pptx"
        dst_pptx = Path(temp_dir) / "rearranged.pptx"

        prs = Presentation()
        # Add 3 slides
        for i in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = f"Slide {i + 1}"
        prs.save(str(src_pptx))

        # Rearrange: [2, 0] (0-indexed: take slide 3, then slide 1)
        res_path = rearrange_slides(src_pptx, sequence=[2, 0], output_path=dst_pptx)
        assert res_path.exists()

        rearranged_prs = Presentation(str(res_path))
        assert len(rearranged_prs.slides) == 2
        assert rearranged_prs.slides[0].shapes.title.text == "Slide 3"
        assert rearranged_prs.slides[1].shapes.title.text == "Slide 1"


def test_pptx_thumbnail_mock():
    """Test generate_thumbnails handles image conversion mock cleanly."""
    with tempfile.TemporaryDirectory() as temp_dir:
        src_pptx = Path(temp_dir) / "test.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        prs.save(str(src_pptx))

        out_prefix = Path(temp_dir) / "thumbs"
        with patch("ccba_ooxml.pptx.thumbnail.convert_to_images", return_value=[]):
            thumbs = generate_thumbnails(src_pptx, out_prefix)
            assert thumbs == []


def test_pptx_replace_text_in_table_and_group():
    """Verify pptx_replace_text handles table cells and grouped shapes."""
    with tempfile.TemporaryDirectory() as temp_dir:
        src_pptx = Path(temp_dir) / "table_group.pptx"
        dst_pptx = Path(temp_dir) / "replaced_table.pptx"

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # blank or title only
        # Add table
        table_shape = slide.shapes.add_table(2, 2, 0, 0, 4000000, 2000000)
        table_shape.table.cell(0, 0).text = "Tiêu đề: {{project_name}}"
        table_shape.table.cell(1, 1).text = "Địa điểm: {{location}}"

        prs.save(str(src_pptx))

        replacements = {
            "{{project_name}}": "Trung tâm Đổi mới Sáng tạo",
            "{{location}}": "Hà Nội",
        }
        res_path = pptx_replace_text(src_pptx, replacements, output_path=dst_pptx)
        assert res_path.exists()

        res_prs = Presentation(str(res_path))
        s0 = res_prs.slides[0]
        found_table = False
        for sp in s0.shapes:
            if sp.has_table:
                found_table = True
                assert "Trung tâm Đổi mới Sáng tạo" in sp.table.cell(0, 0).text
                assert "Hà Nội" in sp.table.cell(1, 1).text
        assert found_table


def test_format_docx_chapter_cover_page():
    """Verify format_docx identifies chapter cover pages and applies administrative typography."""
    with tempfile.TemporaryDirectory() as temp_dir:
        src_docx = Path(temp_dir) / "chapter_doc.docx"
        out_docx = Path(temp_dir) / "formatted_chapter.docx"

        doc = Document()
        doc.add_paragraph("CHƯƠNG I")
        doc.add_paragraph("QUY ĐỊNH CHUNG")
        doc.add_paragraph("Nội dung điều 1 về phạm vi điều chỉnh.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Hạng mục"
        table.cell(0, 1).text = "Giá trị"
        table.cell(1, 0).text = "PCCC"
        table.cell(1, 1).text = "Đạt"
        doc.save(str(src_docx))

        res_path = format_docx(src_docx, out_docx)
        assert res_path.exists()

        formatted_doc = Document(str(res_path))
        ch_para = formatted_doc.paragraphs[0]
        assert ch_para.text == "CHƯƠNG I"
        assert ch_para.runs[0].font.size.pt == 32
        title_para = formatted_doc.paragraphs[1]
        assert title_para.text == "QUY ĐỊNH CHUNG"
        assert title_para.runs[0].font.size.pt == 28
