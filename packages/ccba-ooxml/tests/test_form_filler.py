# packages/ccba-ooxml/tests/test_form_filler.py
"""Unit tests for Word Form Filler and Form Layout Guard (Issue #296)."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import docx
import pytest
from docx.oxml.ns import qn

from ccba_ooxml.form_filler import (
    FormFillConfig,
    FormLayoutGuard,
    SofficeFallbackEngine,
    TableRule,
    TemplateNotFoundError,
    TemplateProtectionError,
    WinwordEngine,
    WordFormFiller,
)

pytestmark = [pytest.mark.fast, pytest.mark.unit]


def test_models_validation():
    """Validates FormFillConfig and TableRule schemas."""
    config = FormFillConfig(
        engine="auto",
        keep_font_formatting=True,
        prevent_row_split=True,
        prune_empty_rows=True,
        page_break_keywords=["PHẦN 2", "MỤC III"],
    )
    assert config.engine == "auto"
    assert config.keep_font_formatting is True
    assert len(config.page_break_keywords) == 2

    rule = TableRule(
        table_index=1,
        header_keyword="DANH SÁCH",
        data_rows=[{"Họ tên": "Nguyen Van A", "Năm sinh": "1990"}],
        delete_unused_template_rows=True,
        allow_break_across_pages=False,
    )
    assert rule.table_index == 1
    assert len(rule.data_rows) == 1
    assert rule.allow_break_across_pages is False


def test_template_not_found(tmp_path: Path):
    """Verifies that non-existent template raises TemplateNotFoundError."""
    missing_file = tmp_path / "missing_template.doc"
    with pytest.raises(TemplateNotFoundError):
        WordFormFiller(missing_file)


def test_layout_guard_docx_cant_split():
    """Verifies injection of <w:cantSplit/> XML element into table row properties."""
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    row = table.rows[0]

    tr_pr = row._tr.get_or_add_trPr()
    assert tr_pr.find(qn("w:cantSplit")) is None

    FormLayoutGuard.apply_docx_cant_split(row)
    assert tr_pr.find(qn("w:cantSplit")) is not None

    # Calling twice should be idempotent
    FormLayoutGuard.apply_docx_cant_split(row)
    assert len(tr_pr.findall(qn("w:cantSplit"))) == 1


def test_layout_guard_empty_row_detection_and_removal():
    """Verifies detection and clean removal of empty table rows."""
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"

    table.rows[1].cells[0].text = "Filled Data"
    table.rows[1].cells[1].text = "Some Value"

    # Row 2 remains empty
    assert not FormLayoutGuard.is_docx_row_empty(table.rows[0])
    assert not FormLayoutGuard.is_docx_row_empty(table.rows[1])
    assert FormLayoutGuard.is_docx_row_empty(table.rows[2])

    FormLayoutGuard.remove_docx_row(table, table.rows[2])
    assert len(table.rows) == 2


def test_fallback_engine_paragraph_and_table_filling(tmp_path: Path):
    """Verifies end-to-end template population using python-docx fallback engine."""
    # 1. Create a dummy DOCX template
    template_path = tmp_path / "test_template.docx"
    doc = docx.Document()
    doc.add_paragraph("HỒ SƠ THỊ THỰC / VISA APPLICATION")
    doc.add_paragraph("HỌ VÀ TÊN: {{HO_VA_TEN}}")
    doc.add_paragraph("NGÀY SINH: {{NGAY_SINH}}")

    # Add dynamic table: row 0 is header, row 1 is template row, row 2 is unused empty row
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "STT"
    table.rows[0].cells[1].text = "Họ Tên"
    table.rows[0].cells[2].text = "Quan Hệ"

    table.rows[1].cells[0].text = "{{STT}}"
    table.rows[1].cells[1].text = "{{HO_TEN}}"
    table.rows[1].cells[2].text = "{{QUAN_HE}}"

    doc.add_paragraph("PHẦN KẾT LUẬN: Đã kiểm tra đầy đủ hồ sơ.")
    doc.save(str(template_path))

    # 2. Run WordFormFiller in soffice/fallback mode
    out_docx = tmp_path / "filled_output.docx"
    config = FormFillConfig(
        engine="soffice",
        prevent_row_split=True,
        prune_empty_rows=True,
        page_break_keywords=["PHẦN KẾT LUẬN"],
    )

    with WordFormFiller(template_path, config=config) as filler:
        filler.apply_paragraphs(
            {
                "{{HO_VA_TEN}}": "VŨ VĂN CHỦ",
                "{{NGAY_SINH}}": "15/08/1990",
            }
        )
        filler.apply_tables(
            [
                TableRule(
                    table_index=0,
                    data_rows=[
                        {"STT": "1", "Họ Tên": "Nguyễn Văn A", "Quan Hệ": "Bố"},
                        {"STT": "2", "Họ Tên": "Trần Thị B", "Quan Hệ": "Mẹ"},
                    ],
                    delete_unused_template_rows=True,
                    allow_break_across_pages=False,
                )
            ]
        )
        filler.export(doc_out=out_docx)

    assert out_docx.exists()

    # 3. Inspect generated document
    result_doc = docx.Document(str(out_docx))
    full_text = "\n".join(p.text for p in result_doc.paragraphs)
    assert "VŨ VĂN CHỦ" in full_text
    assert "15/08/1990" in full_text
    assert "{{HO_VA_TEN}}" not in full_text

    # Verify table rows
    res_table = result_doc.tables[0]
    # Header + 2 data rows = 3 rows total (empty row 2 pruned)
    assert len(res_table.rows) == 3
    assert res_table.rows[1].cells[1].text == "Nguyễn Văn A"
    assert res_table.rows[2].cells[1].text == "Trần Thị B"

    # Verify cantSplit guard applied across all rows
    for r in res_table.rows:
        tr_pr = r._tr.get_or_add_trPr()
        assert tr_pr.find(qn("w:cantSplit")) is not None


def test_winword_engine_mocked(tmp_path: Path):
    """Verifies WinwordEngine execution flow with a mocked Word COM application."""
    template_path = tmp_path / "mock_template.docx"
    template_path.write_text("Dummy binary content", encoding="utf-8")

    # Mock Word COM DOM hierarchy
    mock_word = MagicMock()
    mock_doc = MagicMock()
    mock_word.Documents.Open.return_value = mock_doc

    # Mock Paragraph
    mock_p = MagicMock()
    mock_p.Range.Text = "HỌ VÀ TÊN: {{HO_VA_TEN}}"
    mock_doc.Paragraphs = [mock_p]

    # Mock Table
    mock_table = MagicMock()
    mock_doc.Tables = MagicMock()
    mock_doc.Tables.Count = 1
    mock_doc.Tables.Item.return_value = mock_table
    mock_doc.Tables.__iter__.return_value = [mock_table]

    mock_row = MagicMock()
    mock_table.Rows = MagicMock()
    mock_table.Rows.Count = 1
    mock_table.Rows.__iter__.return_value = [mock_row]

    engine = WinwordEngine(
        template_path=template_path,
        config=FormFillConfig(prevent_row_split=True),
        word_app=mock_word,
    )

    engine.apply_paragraphs({"{{HO_VA_TEN}}": "TEST USER"})
    assert mock_p.Range.Text == "HỌ VÀ TÊN: TEST USER"

    engine.apply_layout_guard()
    assert mock_row.AllowBreakAcrossPages is False

    out_doc = tmp_path / "mock_out.doc"
    out_pdf = tmp_path / "mock_out.pdf"
    engine.export(doc_out=out_doc, pdf_out=out_pdf)

    assert mock_doc.SaveAs2.called
    assert mock_doc.ExportAsFixedFormat.called

    engine.close()
    mock_doc.Close.assert_called_with(SaveChanges=False)


def test_engine_auto_detection_on_linux(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """Verifies that engine='auto' resolves to SofficeFallbackEngine on non-Windows platforms."""
    monkeypatch.setattr("ccba_ooxml.form_filler._is_windows", lambda: False)
    template_path = tmp_path / "auto_test.docx"
    doc = docx.Document()
    doc.add_paragraph("Sample")
    doc.save(str(template_path))

    with WordFormFiller(template_path, engine="auto") as filler:
        assert isinstance(filler._engine, SofficeFallbackEngine)


def test_template_protection_error_raises(tmp_path: Path):
    """Verifies that attempting to export to template path raises TemplateProtectionError."""
    template_path = tmp_path / "template.docx"
    doc = docx.Document()
    doc.add_paragraph("Template")
    doc.save(str(template_path))

    with WordFormFiller(template_path, engine="soffice") as filler:
        with pytest.raises(TemplateProtectionError):
            filler.export(doc_out=template_path)

        with pytest.raises(TemplateProtectionError):
            filler.export(pdf_out=template_path)


def test_auto_map_fields_paragraphs_and_checkboxes(tmp_path: Path):
    """Verifies auto_map_fields correctly handles inline labels with dots, placeholders, and checkboxes."""
    template_path = tmp_path / "form_template.docx"
    doc = docx.Document()
    doc.add_paragraph("ĐƠN XIN CẤP THỊ THỰC")
    p1 = doc.add_paragraph()
    r1 = p1.add_run("1. Họ và tên: ")
    r1.bold = True
    p1.add_run("........................................")
    doc.add_paragraph("2. Ngày tháng năm sinh: ....................   Số CCCD: {{so_cmnd_cccd}}")
    doc.add_paragraph("3. Giới tính: Nam [ ]   Nữ [ ]")
    doc.add_paragraph("4. Tình trạng hôn nhân: [ ] Đã kết hôn   [ ] Độc thân")
    doc.save(str(template_path))

    data = {
        "ho_va_ten": "VŨ VĂN CHỦ",
        "ngay_sinh": "15/08/1990",
        "so_cmnd_cccd": "001090001234",
        "gioi_tinh": "Nam",
        "da_ket_hon": True,
    }

    out_docx = tmp_path / "filled_form.docx"
    with WordFormFiller(template_path, engine="soffice") as filler:
        filler.auto_map_fields(data).export(doc_out=out_docx)

    assert out_docx.exists()
    res_doc = docx.Document(str(out_docx))
    full_text = "\n".join(p.text for p in res_doc.paragraphs)

    assert "VŨ VĂN CHỦ" in full_text
    assert "15/08/1990" in full_text
    assert "001090001234" in full_text
    assert "Nam [X]" in full_text
    assert "Nữ [ ]" in full_text
    assert "[X] Đã kết hôn" in full_text
    assert "[ ] Độc thân" in full_text


def test_auto_map_fields_tables_and_lists(tmp_path: Path):
    """Verifies auto_map_fields fills property sheet tables and dynamic data tables."""
    template_path = tmp_path / "tables_template.docx"
    doc = docx.Document()

    # Table 0: Property Sheet (4-column key-value)
    t0 = doc.add_table(rows=2, cols=4)
    t0.rows[0].cells[0].text = "Họ và tên:"
    t0.rows[0].cells[1].text = "...................."
    t0.rows[0].cells[2].text = "Quốc tịch:"
    t0.rows[0].cells[3].text = ""

    t0.rows[1].cells[0].text = "Nơi ở hiện nay:"
    t0.rows[1].cells[1].text = "...................."
    t0.rows[1].cells[2].text = "Số điện thoại:"
    t0.rows[1].cells[3].text = "{{so_dien_thoai}}"

    # Table 1: Dynamic Data Table (list of dicts)
    t1 = doc.add_table(rows=2, cols=4)
    t1.rows[0].cells[0].text = "STT"
    t1.rows[0].cells[1].text = "Họ và tên"
    t1.rows[0].cells[2].text = "Quan hệ"
    t1.rows[0].cells[3].text = "Năm sinh"

    t1.rows[1].cells[0].text = "{{STT}}"
    t1.rows[1].cells[1].text = "{{HO_TEN}}"
    t1.rows[1].cells[2].text = "{{QUAN_HE}}"
    t1.rows[1].cells[3].text = "{{NAM_SINH}}"

    doc.save(str(template_path))

    data = {
        "ho_va_ten": "VŨ VĂN CHỦ",
        "quoc_tich": "Việt Nam",
        "noi_o_hien_nay": "Hà Nội",
        "so_dien_thoai": "0901234567",
        "thanh_vien_gia_dinh": [
            {"ho_va_ten": "Nguyễn Văn B", "quan_he": "Bố", "nam_sinh": "1960"},
            {"ho_va_ten": "Trần Thị C", "quan_he": "Mẹ", "nam_sinh": "1965"},
        ],
    }

    out_docx = tmp_path / "filled_tables.docx"
    with WordFormFiller(template_path, engine="soffice") as filler:
        filler.auto_map_fields(data).export(doc_out=out_docx)

    assert out_docx.exists()
    res_doc = docx.Document(str(out_docx))

    # Check Table 0 (Property Sheet)
    t0_res = res_doc.tables[0]
    assert t0_res.rows[0].cells[1].text == "VŨ VĂN CHỦ"
    assert t0_res.rows[0].cells[3].text == "Việt Nam"
    assert t0_res.rows[1].cells[1].text == "Hà Nội"
    assert t0_res.rows[1].cells[3].text == "0901234567"

    # Check Table 1 (Dynamic List Table)
    t1_res = res_doc.tables[1]
    assert len(t1_res.rows) == 3  # Header + 2 data rows
    assert t1_res.rows[1].cells[0].text == "1"
    assert t1_res.rows[1].cells[1].text == "Nguyễn Văn B"
    assert t1_res.rows[1].cells[2].text == "Bố"
    assert t1_res.rows[2].cells[0].text == "2"
    assert t1_res.rows[2].cells[1].text == "Trần Thị C"
    assert t1_res.rows[2].cells[2].text == "Mẹ"


def test_auto_map_fields_multi_inline_partial(tmp_path: Path):
    """Verifies that multiple dot fields on the same line are not cross-contaminated when one is missing."""
    template_path = tmp_path / "multi_inline.docx"
    doc = docx.Document()
    p = doc.add_paragraph()
    r1 = p.add_run("1. Họ và tên: ")
    r1.bold = True
    p.add_run("..........")
    r2 = p.add_run("   2. Quê quán: ")
    r2.bold = True
    p.add_run("..........")
    r3 = p.add_run("   3. Nơi ở: ")
    r3.bold = True
    p.add_run("..........")
    doc.save(str(template_path))

    # Profile only has que_quan and noi_o_hien_nay, omitting ho_va_ten
    data = {
        "que_quan": "Nghệ An",
        "noi_o_hien_nay": "Hà Nội",
    }

    out_docx = tmp_path / "multi_inline_out.docx"
    with WordFormFiller(template_path, engine="soffice") as filler:
        filler.auto_map_fields(data).export(doc_out=out_docx)

    res_doc = docx.Document(str(out_docx))
    text = res_doc.paragraphs[0].text
    # Họ và tên dots must remain untouched, NOT replaced by que_quan
    assert "1. Họ và tên: .........." in text
    assert "2. Quê quán: Nghệ An" in text
    assert "3. Nơi ở: Hà Nội" in text

    # Verify bold formatting preserved on labels
    runs = res_doc.paragraphs[0].runs
    assert runs[0].bold is True  # Label 1
    assert runs[2].bold is True  # Label 2
    assert runs[4].bold is True  # Label 3


def test_auto_map_fields_checkboxes_context_and_spacing(tmp_path: Path):
    """Verifies checkbox matching is context-aware and handles multiple spaces between label and box."""
    template_path = tmp_path / "cb_context.docx"
    doc = docx.Document()
    doc.add_paragraph("Giới tính: Nam  [ ]   Nữ  [ ]")
    doc.add_paragraph("Tình trạng: [ ]  Đã kết hôn   [ ]  Độc thân")
    doc.save(str(template_path))

    # User is female (Nữ) whose first name happens to be 'Nam'
    data = {
        "ten": "Nam",
        "gioi_tinh": "Nữ",
        "da_ket_hon": True,
    }

    out_docx = tmp_path / "cb_context_out.docx"
    with WordFormFiller(template_path, engine="soffice") as filler:
        filler.auto_map_fields(data).export(doc_out=out_docx)

    res_doc = docx.Document(str(out_docx))
    p0 = res_doc.paragraphs[0].text
    p1 = res_doc.paragraphs[1].text

    assert "Nam  [ ]" in p0  # Must NOT be checked despite ten == 'Nam'
    assert "Nữ  [X]" in p0
    assert "[X]  Đã kết hôn" in p1
    assert "[ ]  Độc thân" in p1


def test_auto_map_fields_dynamic_table_with_stt_and_alias(tmp_path: Path):
    """Verifies dynamic tables correctly populate STT column even when STT is present in data dict."""
    template_path = tmp_path / "stt_alias.docx"
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text = "STT"
    t.rows[0].cells[1].text = "Họ và tên"
    t.rows[0].cells[2].text = "Quan hệ"

    t.rows[1].cells[0].text = "{{STT}}"
    t.rows[1].cells[1].text = "{{HO_TEN}}"
    t.rows[1].cells[2].text = "{{QUAN_HE}}"
    doc.save(str(template_path))

    data = {
        "thanh_vien_gia_dinh": [
            {"stt": 1, "ho_ten": "Nguyễn Văn A", "quan_he": "Bố"},
            {"stt": 2, "ho_ten": "Trần Thị B", "quan_he": "Mẹ"},
        ]
    }

    out_docx = tmp_path / "stt_alias_out.docx"
    with WordFormFiller(template_path, engine="soffice") as filler:
        filler.auto_map_fields(data).export(doc_out=out_docx)

    res_doc = docx.Document(str(out_docx))
    res_t = res_doc.tables[0]
    assert res_t.rows[1].cells[0].text == "1"
    assert res_t.rows[1].cells[1].text == "Nguyễn Văn A"
    assert res_t.rows[1].cells[2].text == "Bố"
    assert res_t.rows[2].cells[0].text == "2"
    assert res_t.rows[2].cells[1].text == "Trần Thị B"
    assert res_t.rows[2].cells[2].text == "Mẹ"


def test_normalize_label_and_que_quan():
    """Validates label normalization with prefixes and distinct que_quan mapping."""
    from ccba_ooxml.form_filler.aliases import normalize_label, resolve_field_value

    assert normalize_label("1. Họ và tên:") == "ho_va_ten"
    assert normalize_label("1.1. Họ và tên:") == "ho_va_ten"
    assert normalize_label("a. Họ và tên:") == "ho_va_ten"
    assert normalize_label("a) Họ và tên:") == "ho_va_ten"
    assert normalize_label("(1) Họ và tên:") == "ho_va_ten"
    assert normalize_label("I. Họ và tên:") == "ho_va_ten"

    # Verify que_quan is distinct from noi_sinh
    data = {"noi_sinh": "Hà Nội", "que_quan": "Nghệ An"}
    f1, v1, k1 = resolve_field_value("Quê quán", data)
    f2, v2, k2 = resolve_field_value("Nơi sinh", data)
    assert f1 and v1 == "Nghệ An" and k1 == "que_quan"
    assert f2 and v2 == "Hà Nội" and k2 == "noi_sinh"
