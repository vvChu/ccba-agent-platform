"""Comprehensive Scoped Unit Tests for CCBA PPTX Deck Builder & Parser."""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation

from ccba_ooxml import (
    CardItem,
    CCBAPresentationTheme,
    DeckBuilder,
    MarkdownDeckParser,
    SlideSpec,
    SlideType,
    build_presentation_from_markdown,
    unpack_document,
)
from ccba_ooxml.pptx.templates import (
    COLOR_CCBA_DARK_BLUE,
    COLOR_IBST_RED,
    get_logo_asset_path,
)


@pytest.mark.fast
def test_markdown_parser_frontmatter_and_sections() -> None:
    """Test AST parsing from YAML frontmatter and standard sections."""
    md = """---
title: "Báo Cáo Nghiên Cứu BIM"
subtitle: "Ứng dụng trong Thẩm tra Thiết kế"
author: "KTS. Nguyễn Văn A"
date: "02/09/2026"
---

# [QUY CHUẨN] Tổng Quan Nghị Định 217/2026
### Quy định bắt buộc áp dụng BIM
- Quy định lộ trình áp dụng BIM theo Luật Xây dựng 2025.
- Yêu cầu bàn giao mô hình định dạng IFC và CDE chuẩn ISO 19650.
- Trách nhiệm của Chủ đầu tư và Tư vấn Thẩm tra.

<!-- notes: Diễn giải kỹ về thời hạn áp dụng từ 01/07/2026 -->
"""
    specs = MarkdownDeckParser.parse_markdown(md)
    assert len(specs) == 2

    # Cover Slide
    cover = specs[0]
    assert cover.slide_type == SlideType.COVER
    assert cover.title == "Báo Cáo Nghiên Cứu BIM"
    assert cover.subtitle == "Ứng dụng trong Thẩm tra Thiết kế"
    assert cover.author == "KTS. Nguyễn Văn A"
    assert cover.date == "02/09/2026"

    # Content Slide
    slide1 = specs[1]
    assert slide1.slide_type == SlideType.STANDARD
    assert slide1.badge == "QUY CHUẨN"
    assert "Tổng Quan Nghị Định 217/2026" in slide1.title
    assert slide1.subtitle == "Quy định bắt buộc áp dụng BIM"
    assert len(slide1.bullets) == 3
    assert "Diễn giải kỹ" in slide1.notes


@pytest.mark.fast
def test_markdown_parser_split_and_tables() -> None:
    """Test parsing of 2-column split layouts and native markdown tables."""
    md = """
# [SO SÁNH] Quy Trình Truyền Thống vs CCBA WAY

::: split
### Quy Trình Cũ
- Bản vẽ 2D rời rạc
- Xung đột phát hiện tại hiện trường
- Chi phí phát sinh cao

::: col-2
### CCBA WAY
- Mô hình BIM 3D đồng bộ
- Thẩm tra phát hiện xung đột sớm
- Tối ưu hóa thời gian và chi phí
:::

---

# [BẢNG SỐ LIỆU] Thống Kê Xung Đột Mô Hình

| Bộ môn | Số lượng Va chạm | Mức độ Nghiêm trọng | Trạng thái |
| :--- | :--- | :--- | :--- |
| Kiến trúc - Kết cấu | 12 | Cao (Critical) | Đang xử lý |
| Kết cấu - MEP | 45 | Trung bình (Medium) | Đã gán việc |
| MEP - PCCC | 8 | Nghiêm trọng | Đã khắc phục |
"""
    specs = MarkdownDeckParser.parse_markdown(md)
    assert len(specs) == 2

    # Split slide
    split_slide = specs[0]
    assert split_slide.slide_type == SlideType.SPLIT
    assert split_slide.left_title == "Quy Trình Cũ"
    assert split_slide.right_title == "CCBA WAY"
    assert len(split_slide.left_column) == 3
    assert len(split_slide.right_column) == 3

    # Table slide
    table_slide = specs[1]
    assert table_slide.slide_type == SlideType.TABLE
    assert table_slide.table is not None
    assert len(table_slide.table.headers) == 4
    assert len(table_slide.table.rows) == 3
    assert table_slide.table.headers[0] == "Bộ môn"
    assert table_slide.table.rows[0][0] == "Kiến trúc - Kết cấu"


@pytest.mark.fast
def test_markdown_parser_callouts_and_cards() -> None:
    """Test parsing of GitHub-style Alert blocks into CardItem objects."""
    md = """
# [LƯU Ý] Các Điểm Cần Chú Ý Trong Thiết Kế PCCC

> [!NOTE] Căn cứ Pháp lý
> Áp dụng QCVN 06:2022/BXD và Sửa đổi 1:2023.
> Kiểm tra kỹ bậc chịu lửa của công trình.

> [!WARNING] Rủi ro Cốt lõi
> Chiều rộng lối thoát nạn cầu thang bộ không đạt chuẩn.
> Khoảng cách di chuyển đến buồng thang vượt quá quy định.

> [!TIP] Giải pháp Tối ưu
> Tích hợp van ngăn khói tự động trên đường ống MEP xuyên tường.
"""
    specs = MarkdownDeckParser.parse_markdown(md)
    assert len(specs) == 1

    slide = specs[0]
    assert slide.slide_type == SlideType.CARDS
    assert len(slide.cards) == 3
    assert slide.cards[0].title == "Căn cứ Pháp lý"
    assert slide.cards[0].kind == "info"
    assert slide.cards[1].title == "Rủi ro Cốt lõi"
    assert slide.cards[1].kind == "warning"
    assert slide.cards[2].title == "Giải pháp Tối ưu"
    assert slide.cards[2].kind == "success"


@pytest.mark.fast
def test_logo_asset_discovery() -> None:
    """Test that bundled official CCBA logo assets are properly discovered."""
    full_logo = get_logo_asset_path("logo_ccba_full.png")
    assert full_logo is not None
    assert full_logo.exists()


@pytest.mark.fast
def test_deck_builder_compilation(tmp_path: Path) -> None:
    """Test end-to-end PowerPoint presentation compilation with all slide types."""
    output_pptx = tmp_path / "test_presentation.pptx"

    theme = CCBAPresentationTheme(
        primary_color=COLOR_CCBA_DARK_BLUE,
        accent_color=COLOR_IBST_RED,
        aspect_ratio="16:9",
        font_family="Segoe UI",
    )

    specs = [
        SlideSpec(
            title="HỘI THẢO CHUYÊN ĐỀ BIM CCBA",
            subtitle="Áp Dụng Công Nghệ Số Trong Thẩm Tra Kỹ Thuật",
            slide_type=SlideType.COVER,
            author="Ban Kỹ Thuật CCBA",
            date="2026-09-02",
        ),
        SlideSpec(
            title="Tổng Quan Vấn Đề",
            subtitle="Thách thức trong quản lý dự án",
            slide_type=SlideType.STANDARD,
            badge="TỔNG QUAN",
            bullets=[
                "Thiếu sự phối hợp đồng bộ giữa các bộ môn.",
                "Khối lượng bóc tách thủ công dễ sai sót.",
                "Tiến độ thi công bị kéo dài do xử lý va chạm muộn.",
            ],
            notes="Nhấn mạnh vào số liệu thực tế dự án vừa qua.",
        ),
        SlideSpec(
            title="So Sánh Phương Pháp",
            slide_type=SlideType.SPLIT,
            badge="PHÂN TÍCH",
            left_title="2D CAD",
            left_column=["Rời rạc", "Tốn nhân lực", "Rủi ro cao"],
            right_title="3D BIM (CCBA WAY)",
            right_column=["Tập trung", "Tự động hóa", "Kiểm soát rủi ro"],
        ),
        SlideSpec(
            title="Thẻ Ghi Chú & Khuyến Nghị",
            slide_type=SlideType.CARDS,
            badge="KHUYẾN NGHỊ",
            cards=[
                CardItem("Kiến trúc (ARC)", "Rà soát chiều cao thông thủy", "arch"),
                CardItem("Kết cấu (STR)", "Kiểm toán độ võng dầm chuyển", "struct"),
                CardItem("Cơ điện (MEP)", "Tránh xung đột ống gió và dầm", "mep"),
            ],
        ),
        SlideSpec(
            title="Kế Hoạch Triển Khai",
            slide_type=SlideType.SECTION,
            subtitle="Giai đoạn 2: Bàn giao & Phối hợp thông tin",
        ),
    ]

    builder = DeckBuilder(theme=theme)
    out_file = builder.build_from_specs(specs, output_pptx)

    assert out_file.exists()
    assert out_file.stat().st_size > 0

    # Verify with python-pptx
    prs = Presentation(str(out_file))
    assert len(prs.slides) == 5
    assert prs.slide_width == theme.width
    assert prs.slide_height == theme.height


@pytest.mark.fast
def test_build_presentation_from_markdown_helper(tmp_path: Path) -> None:
    """Test the high-level build_presentation_from_markdown convenience function."""
    md_content = """---
title: "Hướng Dẫn Thẩm Tra PCCC"
author: "CCBA Audit Team"
date: "02/09/2026"
---

# [PCCC] Tiêu Chuẩn Áp Dụng
- Áp dụng QCVN 06:2022/BXD.
- Tuân thủ Thông tư 149/2020/TT-BCA.
- Thẩm tra áp lực nước chữa cháy tự động.

---

# [BẢNG] Danh Mục Hồ Sơ Nghiệm Thu

| STT | Tên Hồ sơ | Tình trạng | Ghi chú |
| :--- | :--- | :--- | :--- |
| 01 | Giấy chứng nhận PCCC | Đã duyệt | Hợp lệ |
| 02 | Biên bản thử áp lực đường ống | Hoàn thành | Đạt 10 bar |
"""
    md_file = tmp_path / "seminar_input.md"
    md_file.write_text(md_content, encoding="utf-8")

    out_pptx = tmp_path / "seminar_output.pptx"
    generated_file = build_presentation_from_markdown(md_file, out_pptx)

    assert generated_file.exists()
    assert generated_file.stat().st_size > 0

    prs = Presentation(str(generated_file))
    assert len(prs.slides) == 3


@pytest.mark.fast
def test_pptx_validation_cycle(tmp_path: Path) -> None:
    """Verify generated presentation passes unpack and PPTX validation checks."""
    md_content = """---
title: "Báo Cáo Kiểm Định PPTX"
author: "Quality Gate"
---

# Slide 1: Nội Dung Chính
- Điểm 1
- Điểm 2
"""
    pptx_path = tmp_path / "val_test.pptx"
    build_presentation_from_markdown(md_content, pptx_path)

    unpack_dir = tmp_path / "unpacked_val"
    unpack_document(str(pptx_path), str(unpack_dir))
    assert unpack_dir.exists()
    assert (unpack_dir / "ppt").exists()

    from ccba_ooxml.validation import PPTXSchemaValidator

    # Validate unpacked XML structure against PPTX schemas
    validator = PPTXSchemaValidator(unpack_dir, pptx_path)
    is_valid = validator.validate()
    assert is_valid


@pytest.mark.fast
def test_cli_build_deck_execution(tmp_path: Path) -> None:
    """Test CLI subcommand 'build-deck' via subprocess."""
    md_content = """---
title: "CLI Presentation Test"
author: "CLI Agent"
---

# CLI Slide 1
- Running from CLI
"""
    md_file = tmp_path / "cli_test.md"
    md_file.write_text(md_content, encoding="utf-8")
    out_pptx = tmp_path / "cli_out.pptx"

    cmd = [
        sys.executable,
        "-m",
        "ccba_ooxml",
        "build-deck",
        str(md_file),
        "--output",
        str(out_pptx),
        "--aspect-ratio",
        "16:9",
    ]

    env = os.environ.copy()
    src_dir = Path(__file__).resolve().parent.parent / "src"
    existing_pp = env.get("PYTHONPATH", "")
    env["PYTHONPATH"] = f"{src_dir}{os.pathsep}{existing_pp}" if existing_pp else str(src_dir)

    result = subprocess.run(cmd, capture_output=True, text=True, check=True, env=env)
    assert "SUCCESS" in result.stdout
    assert out_pptx.exists()
    assert out_pptx.stat().st_size > 0


@pytest.mark.fast
def test_markdown_parser_storytelling_archetypes(tmp_path: Path) -> None:
    """Test AST parsing and PowerPoint generation of Cole Knaflic Storytelling archetypes."""
    md = """---
title: "Storytelling Test Deck"
---

# [CHIẾN LƯỢC] Thông Điệp Trọng Tâm

::: big-idea
Kiểm soát mô hình BIM đa chiều là điều kiện tiên quyết để triệt tiêu 100% va chạm trước khi đổ bê tông.
:::
- Bối cảnh: Luật Xây dựng 2025 & Nghị định 217/2026/NĐ-CP
- Rủi ro: Phát sinh chi phí đục phá dầm
- Hành động: Triển khai Quad-View Vision từ bước A0

---

# [LỘ TRÌNH] 4 Bước Thẩm Tra Kỹ Thuật

::: steps
1. **Thu Thập Dữ Liệu**: Chuẩn hóa tệp IFC và CDE
2. **Thẩm Tra Đa Bộ Môn**: Quét 100% va chạm hình học
3. **Lập Báo Cáo Heatmap**: Xếp hạng rủi ro Critical
4. **Cấp Chứng Thư**: Bàn giao cho Chủ đầu tư
:::

---

# [CHƯƠNG TRÌNH] Khung Thảo Luận

::: agenda active=2
- 01. Căn cứ Pháp lý NĐ 217/2026
- 02. Quy trình Thẩm tra CCBA WAY
- 03. Phân tích Dữ liệu Va chạm
- 04. Kế hoạch Nghiệm thu
:::

---

# [Ý KIẾN HIỆN TRƯỜNG] Nhận Xét Từ Ban QLDA

::: quote
Nhờ phát hiện sớm 64 điểm giao cắt giữa dầm và ống gió tại Tầng hầm, dự án đã tiết kiệm hơn 450 triệu đồng.
👤 Ông Trần Văn B — Giám đốc Ban QLDA
:::
"""
    specs = MarkdownDeckParser.parse_markdown(md)
    assert len(specs) == 5

    # 1. Cover
    assert specs[0].slide_type == SlideType.COVER

    # 2. Big Idea
    assert specs[1].slide_type == SlideType.BIG_IDEA
    assert "điều kiện tiên quyết" in specs[1].quote_text
    assert len(specs[1].bullets) == 3

    # 3. Steps
    assert specs[2].slide_type == SlideType.STEPS
    assert len(specs[2].steps) == 4
    assert "Thu Thập Dữ Liệu" in specs[2].steps[0]

    # 4. Agenda
    assert specs[3].slide_type == SlideType.AGENDA
    assert specs[3].active_step == 2
    assert len(specs[3].steps) == 4

    # 5. Quote
    assert specs[4].slide_type == SlideType.QUOTE
    assert "tiết kiệm hơn 450 triệu" in specs[4].quote_text
    assert "Ông Trần Văn B" in specs[4].quote_author

    # Verify building into presentation
    out_file = tmp_path / "storytelling_deck.pptx"
    builder = DeckBuilder()
    generated = builder.build_from_specs(specs, out_file)

    assert generated.exists()
    prs = Presentation(str(generated))
    assert len(prs.slides) == 5
